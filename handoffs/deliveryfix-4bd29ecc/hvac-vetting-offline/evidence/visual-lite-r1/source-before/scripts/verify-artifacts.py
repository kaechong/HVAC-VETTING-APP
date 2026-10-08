"""Independent readback with openpyxl + standard-library XML/ZIP; no network."""
import hashlib
import json
import pathlib
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[3]
MOTHER = ROOT / 'baseline/plan/templates/HVAC_VETTING_CHECKLIST_MOTHER_20261007.xlsx'
HEADERS = ['編號', '複核文件名稱', '問題內容', '違反法規']
sha = lambda data: hashlib.sha256(data).hexdigest()
assert sha(MOTHER.read_bytes()) == 'e968805be7b0e57f94c2a3e391c495f25730c55b9800123b15a6bf8900d56efd'
with zipfile.ZipFile(MOTHER) as z:
    original_parts = {p: z.read(p) for p in z.namelist()}
original = openpyxl.load_workbook(MOTHER)
manifest = json.loads((ROOT / 'samples/final-r2/SAMPLE_MANIFEST.json').read_text())
results = []
for sample in manifest['samples']:
    deliveries = [sample] + ([sample['reissue']] if sample.get('reissue') else [])
    for delivery in deliveries:
        path = pathlib.Path(delivery['report'])
        receipt = json.loads(pathlib.Path(delivery['receipt']).read_text())
        data = path.read_bytes()
        assert sha(data) == delivery['report_sha256'] == receipt['report_sha256']
        with zipfile.ZipFile(path) as z:
            parts = {p: z.read(p) for p in z.namelist()}
        assert set(parts) == set(original_parts)
        changed = []
        for name, part in parts.items():
            ET.fromstring(part)  # catches malformed styles/inline XML independently
            if part != original_parts[name]:
                changed.append(name)
        assert set(changed) <= {'xl/worksheets/sheet2.xml', 'xl/styles.xml', 'xl/sharedStrings.xml'}
        assert parts['xl/worksheets/sheet1.xml'] == original_parts['xl/worksheets/sheet1.xml']
        w = openpyxl.load_workbook(path, data_only=False)
        assert w.sheetnames == original.sheetnames
        s = w['工作表1']
        assert [s.cell(5,c).value for c in range(1,5)] == HEADERS
        assert set(map(str,s.merged_cells.ranges)) == {'A1:D1','A2:D2','A3:D3'}
        assert s.max_column == 4
        assert [s.column_dimensions[c].width for c in ['B','C','D']] == [22.75,46.88,55.38]
        assert s.cell(3,1).value == receipt['version']
        for r in [1,2,3]:
            c=s.cell(r,1)
            assert c.font.name == 'Arial' and c.font.sz == 14 and c.font.b
            assert c.alignment.horizontal == c.alignment.vertical == 'center'
        count = len(receipt['finding_ids'])
        assert s.max_row == max(20,5+count)
        for r in range(6,s.max_row+1):
            assert s.cell(r,1).data_type == 'n' and s.cell(r,1).number_format == '0.0'
            assert s.cell(r,1).value == r-5
            for c in range(1,5):
                cell=s.cell(r,c)
                assert cell.data_type != 'f' and cell.hyperlink is None
                assert cell.font.name == 'Arial' and cell.fill.patternType is None
                assert cell.alignment.horizontal == 'center'
                for edge in ['left','right','top','bottom']:
                    assert getattr(cell.border,edge).style == 'thin'
                    assert getattr(cell.border,edge).color.rgb == 'FF000000'
            if r < 6+count:
                assert s.cell(r,3).value.endswith('；')
                assert s.cell(r,3).alignment.wrap_text
                assert 15.75 <= s.row_dimensions[r].height <= 409.5
            else:
                assert all(s.cell(r,c).value is None for c in range(2,5))
        assert receipt['provisional']['formal_acceptance'] is False
        assert receipt['drive_receipt'] is None
        assert len(receipt['coverage']) == 24 and all(c['state']=='not_reviewed' for c in receipt['coverage'])
        if sample['case']=='long-18':
            streamed=openpyxl.load_workbook(path,read_only=True)['工作表1']
            rows=list(streamed.iter_rows(values_only=True))
            assert len(rows)==23 and [r[0] for r in rows[20:23]]==[16,17,18]
        results.append({'case':sample['case'],'version':receipt['version'],'status':'PASS_OFFLINE_READBACK','report_sha256':sha(data),'findings':count,'qc':receipt['qc'],'changed_parts':changed})
print(json.dumps({'reader':'openpyxl '+openpyxl.__version__, 'scope':'synthetic workbook structure/content/hash; not visual or engineering acceptance', 'results':results},ensure_ascii=False,indent=2))
